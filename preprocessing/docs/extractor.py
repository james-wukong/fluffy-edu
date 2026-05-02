# processing/docs/extractor.py

import re
from pathlib import Path
from typing import cast

import filetype
from docx import Document
from docx.document import Document as DocumentType
from docx.table import Table

from preprocessing.schema import (
    DocMetadata,
    ExtractHelper,
    FileMetadata,
    MetaData,
    PageResult,
)


class DocxExtractor:
    def extract_file(self, file_path: str) -> list[PageResult]:
        if not file_path.endswith((".docx", ".doc")):
            raise TypeError("File must be a .docx file")

        path_obj = Path(file_path)
        doc = Document(file_path)
        kind = filetype.guess(file_path)
        metadata = cast(MetaData, {})
        file_meta: FileMetadata = {
            "filename": path_obj.name,
            "file_path": path_obj.parent,
            "file_type": kind.extension if kind else "",
        }
        metadata["file_meta"] = file_meta
        doc_meta = DocMetadata()

        # Extract core properties if available
        props = doc.core_properties
        if props:
            doc_meta = ExtractHelper.map_doc_metadata(props)
            metadata["doc_meta"] = doc_meta

        results: list[PageResult] = []
        # 1. Extract body paragraphs (grouped by section/heading)
        results.extend(self._extract_paragraphs(doc, metadata))

        # 2. Extract tables separately
        results.extend(self._extract_tables(doc, metadata))

        return results

    # ── Paragraph extraction ───────────────────────────────────

    def _extract_paragraphs(
        self, doc: DocumentType, metadata: MetaData
    ) -> list[PageResult]:
        """
        Group paragraphs under their nearest heading.
        Each heading + its body paragraphs = one PageResult.
        This keeps context together instead of splitting randomly.
        """
        results: list[PageResult] = []
        # current_heading: str = ""
        # current_level: int = 0
        current_text: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (
                para.style.name
                if para is not None and para.style is not None
                else "Normal"
            )

            if style is not None and style.startswith("Heading"):
                # Save previous section
                if current_text:
                    results.append(
                        {
                            "text": "\n".join(current_text),
                            "metadata": metadata,
                        }
                    )

                # Start new section
                # current_heading = text
                # current_level = self._heading_level(style)
                current_text = [text]

            else:
                current_text.append(text)

        # Save final section
        if current_text:
            results.append(
                {
                    "text": "\n".join(current_text),
                    "metadata": metadata,
                }
            )

        return results

    def _heading_level(self, style_name: str) -> int:
        """Extract level from 'Heading 1' → 1"""
        match = re.search(r"\d+", style_name)
        return int(match.group()) if match else 0

    # ── Table extraction ───────────────────────────────────────

    def _extract_tables(
        self, doc: DocumentType, metadata: MetaData
    ) -> list[PageResult]:
        results = []

        for table in doc.tables:
            text = self._table_to_markdown(table)
            if not text:
                continue

            results.append(
                {
                    "text": text,
                    "metadata": metadata,
                }
            )

        return results

    def _table_to_markdown(self, table: Table) -> str:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(cells)

        if not rows:
            return ""

        # First row = headers
        headers = rows[0]
        separator = ["---"] * len(headers)
        data_rows = rows[1:]

        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(separator) + " |",
        ]
        for row in data_rows:
            lines.append("| " + " | ".join(row) + " |")

        return "\n".join(lines)
